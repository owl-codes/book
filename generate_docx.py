#!/usr/bin/env python3
"""Generate a Word document from the book's markdown chapters."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CHAPTERS_DIR = Path(__file__).parent / "chapters"
OUTPUT = Path(__file__).parent / "The_Wildtongue_Chronicles.docx"

CHAPTER_FILES = [
    "01-the-frequency.md",
    "02-the-laughing-man.md",
    "03-the-dive.md",
    "04-the-lawkeeper.md",
]


def parse_inline(paragraph, text, base_bold=False, base_italic=False):
    """Parse markdown inline formatting (*italic*, **bold**, ***both***) into runs."""
    # Pattern matches ***bold+italic***, **bold**, or *italic*
    pattern = re.compile(r'(\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(.+?)\*)')
    pos = 0
    for m in pattern.finditer(text):
        # Add text before this match
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            run.bold = base_bold
            run.italic = base_italic
        if m.group(2):  # ***bold+italic***
            run = paragraph.add_run(m.group(2))
            run.bold = True
            run.italic = True
        elif m.group(3):  # **bold**
            run = paragraph.add_run(m.group(3))
            run.bold = True
            run.italic = base_italic
        elif m.group(4):  # *italic*
            run = paragraph.add_run(m.group(4))
            run.bold = base_bold
            run.italic = True
        pos = m.end()
    # Remaining text
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = base_bold
        run.italic = base_italic


def build_doc():
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

    # Default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Georgia'
    font.size = Pt(11.5)
    pf = style.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.35

    # Title page
    for _ in range(6):
        doc.add_paragraph('')

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('THE WILDTONGUE CHRONICLES')
    run.bold = True
    run.font.size = Pt(26)
    run.font.name = 'Georgia'

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Book One: The Choir of Erythis-3')
    run.italic = True
    run.font.size = Pt(16)
    run.font.name = 'Georgia'
    run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_page_break()

    # Process each chapter
    for i, filename in enumerate(CHAPTER_FILES):
        filepath = CHAPTERS_DIR / filename
        md = filepath.read_text(encoding='utf-8')
        lines = md.split('\n')

        for line in lines:
            stripped = line.strip()

            # Skip empty lines (we handle spacing via paragraph formatting)
            if not stripped:
                continue

            # Chapter heading (# Chapter ...)
            if stripped.startswith('# '):
                heading_text = stripped[2:]
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(48)
                p.paragraph_format.space_after = Pt(24)
                run = p.add_run(heading_text.upper())
                run.bold = True
                run.font.size = Pt(18)
                run.font.name = 'Georgia'
                continue

            # Scene break (---)
            if stripped == '---':
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(16)
                run = p.add_run('* * *')
                run.font.size = Pt(12)
                run.font.name = 'Georgia'
                run.font.color.rgb = RGBColor(120, 120, 120)
                continue

            # Bold-only lines (like **Erythis-3.**)
            bold_match = re.match(r'^\*\*(.+)\*\*$', stripped)
            if bold_match:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(bold_match.group(1))
                run.bold = True
                continue

            # Normal paragraph with inline formatting
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            parse_inline(p, stripped)

        # Page break between chapters (not after last)
        if i < len(CHAPTER_FILES) - 1:
            doc.add_page_break()

    doc.save(str(OUTPUT))
    print(f"Generated: {OUTPUT}")


if __name__ == '__main__':
    build_doc()
